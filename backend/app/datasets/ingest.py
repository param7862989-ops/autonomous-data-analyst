"""Parse untrusted CSV / XLSX uploads into a table of raw cell values.

Nothing in an uploaded file is executed: workbooks are read with ``data_only=True`` (cached
formula results, never formulas), macro-enabled and legacy/encrypted workbooks are rejected, the
archive is checked for zip bombs before parsing, and XML parsing goes through defusedxml (used by
openpyxl automatically when installed).
"""

from __future__ import annotations

import codecs
import csv
import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any, Literal

import pandas as pd

from app.config import Settings
from app.errors import AppError, PayloadTooLargeError

FileFormat = Literal["csv", "xlsx"]

_ZIP_SIGNATURE = b"PK\x03\x04"
# OLE2 compound file: legacy .xls, and also password-protected .xlsx files.
_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_CSV_EXTENSIONS = frozenset({".csv", ".tsv", ".txt"})
_MAX_ZIP_ENTRIES = 10_000
_MAX_LEADING_EMPTY_ROWS = 1_000
_EXCEL_MAX_COLUMNS = 16_384
_EXCEL_MAX_ROWS = 1_048_576
# Empty rows tolerated beyond the row limit before a sheet is rejected. openpyxl yields an empty
# row for every gap in row numbers, so a tiny file can claim billions of rows.
_EMPTY_ROW_ALLOWANCE = 100_000


class DatasetFileError(AppError):
    """The uploaded file cannot be turned into a dataset. Messages are safe to show users."""

    status_code = 422
    code = "invalid_file"


@dataclass(frozen=True)
class IngestLimits:
    max_rows: int
    max_columns: int
    max_cells: int
    max_xlsx_uncompressed_bytes: int

    @classmethod
    def from_settings(cls, settings: Settings) -> IngestLimits:
        return cls(
            max_rows=settings.max_dataset_rows,
            max_columns=settings.max_dataset_columns,
            max_cells=settings.max_dataset_cells,
            max_xlsx_uncompressed_bytes=settings.max_xlsx_uncompressed_bytes,
        )


@dataclass
class ParseNote:
    code: str
    message: str


@dataclass
class RawTable:
    """Header strings exactly as found, plus a positional frame of raw cell values."""

    headers: list[str]
    frame: pd.DataFrame
    file_format: FileFormat
    sheet_name: str | None = None
    encoding: str | None = None
    delimiter: str | None = None
    notes: list[ParseNote] = field(default_factory=list)


# --- Format detection ----------------------------------------------------------------------


def detect_format(filename: str, head: bytes) -> FileFormat:
    """Decide the format from the extension and confirm it against the file's leading bytes."""
    extension = PurePath(filename).suffix.lower()
    if extension in {".xls", ".xlsm", ".xlsb", ".ods"}:
        raise DatasetFileError(
            f"{extension} files aren't supported. Save the file as .xlsx (Excel workbook) or "
            "CSV and upload it again.",
            code="unsupported_file_type",
        )
    if extension == ".xlsx":
        if head.startswith(_ZIP_SIGNATURE):
            return "xlsx"
        if head.startswith(_OLE_SIGNATURE):
            raise DatasetFileError(
                "This workbook is password-protected or in the legacy .xls format. Remove the "
                "password or save it as a regular .xlsx file.",
                code="unsupported_file_type",
            )
        raise DatasetFileError(
            "The file has an .xlsx extension but isn't an Excel workbook.",
            code="file_type_mismatch",
        )
    if extension in _CSV_EXTENSIONS:
        if head.startswith(_ZIP_SIGNATURE) or head.startswith(_OLE_SIGNATURE):
            raise DatasetFileError(
                "The file has a CSV extension but contains binary data.",
                code="file_type_mismatch",
            )
        return "csv"
    raise DatasetFileError(
        "Upload a .csv or .xlsx file.",
        code="unsupported_file_type",
    )


# --- Limits ----------------------------------------------------------------------------------


def _check_shape(rows: int, columns: int, limits: IngestLimits) -> None:
    if columns > limits.max_columns:
        raise DatasetFileError(
            f"The file has {columns:,} columns; the limit is {limits.max_columns:,}.",
            code="too_many_columns",
        )
    if rows > limits.max_rows:
        raise DatasetFileError(
            f"The file has more than {limits.max_rows:,} rows. Upload a smaller extract.",
            code="too_many_rows",
        )
    if rows * columns > limits.max_cells:
        raise DatasetFileError(
            f"The file has {rows * columns:,} cells; the limit is {limits.max_cells:,}.",
            code="too_many_cells",
        )


# --- CSV -------------------------------------------------------------------------------------


def _decode(data: bytes) -> tuple[str, str, list[ParseNote]]:
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16"), "utf-16", []
    try:
        return data.decode("utf-8-sig"), "utf-8", []
    except UnicodeDecodeError:
        pass
    note = ParseNote(
        code="encoding_fallback",
        message="The file isn't valid UTF-8, so it was read as Windows-1252. Check that accented "
        "characters look right.",
    )
    try:
        return data.decode("cp1252"), "windows-1252", [note]
    except UnicodeDecodeError:
        note.message = note.message.replace("Windows-1252", "Latin-1")
        return data.decode("latin-1"), "latin-1", [note]


def _sniff_delimiter(text: str) -> str:
    sample = text[:65_536]
    if "\n" in sample:
        sample = sample[: sample.rfind("\n")]
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


_BAD_LINE = re.compile(r"Expected (\d+) fields in line (\d+), saw (\d+)")


def _header_width(text: str, delimiter: str) -> int:
    """Number of fields in the first record, counted without building a DataFrame."""
    try:
        header = next(csv.reader(io.StringIO(text), delimiter=delimiter), [])
    except csv.Error as exc:
        raise DatasetFileError(
            "The first row couldn't be read as CSV column names.", code="malformed_file"
        ) from exc
    return len(header)


def read_csv(data: bytes, limits: IngestLimits) -> RawTable:
    if not data.strip():
        raise DatasetFileError("The file is empty.", code="empty_file")
    text, encoding, notes = _decode(data)
    if "\x00" in text:
        raise DatasetFileError(
            "The file contains binary data and doesn't look like a CSV file.",
            code="file_type_mismatch",
        )
    delimiter = _sniff_delimiter(text)
    # Check the width before pandas parses anything: a single line of many thousands of empty
    # fields is tiny on disk but very expensive as a DataFrame.
    _check_shape(0, _header_width(text, delimiter), limits)
    try:
        frame = pd.read_csv(
            io.StringIO(text),
            sep=delimiter,
            header=None,
            dtype=str,
            keep_default_na=False,
            na_filter=False,
            skip_blank_lines=True,
            # Header row + max rows + 1 extra row to detect overflow without reading everything.
            nrows=limits.max_rows + 2,
            on_bad_lines="error",
            engine="c",
        )
    except pd.errors.EmptyDataError as exc:
        raise DatasetFileError("The file is empty.", code="empty_file") from exc
    except pd.errors.ParserError as exc:
        match = _BAD_LINE.search(str(exc))
        message = (
            f"Line {match.group(2)} has {match.group(3)} values but the first row has "
            f"{match.group(1)}. Check for unquoted delimiters in that line."
            if match
            else "The file couldn't be parsed as CSV. Check that it's a valid, delimited file."
        )
        raise DatasetFileError(message, code="malformed_file") from exc

    if frame.shape[0] < 1:
        raise DatasetFileError("The file is empty.", code="empty_file")
    headers = ["" if pd.isna(v) else str(v) for v in frame.iloc[0].tolist()]
    body = frame.iloc[1:].reset_index(drop=True)
    body.columns = list(range(body.shape[1]))
    _check_shape(body.shape[0], body.shape[1], limits)
    if body.shape[0] == 0:
        raise DatasetFileError("The file has column names but no data rows.", code="no_data_rows")
    return RawTable(
        headers=headers,
        frame=body,
        file_format="csv",
        encoding=encoding,
        delimiter=delimiter,
        notes=notes,
    )


# --- XLSX ------------------------------------------------------------------------------------


def _validate_workbook_archive(data: bytes, limits: IngestLimits) -> None:
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise DatasetFileError(
            "The workbook is damaged and couldn't be opened.", code="malformed_file"
        ) from exc
    with archive:
        entries = archive.infolist()
        if len(entries) > _MAX_ZIP_ENTRIES:
            raise DatasetFileError("The workbook has too many parts.", code="malformed_file")
        names = {entry.filename for entry in entries}
        if "xl/workbook.xml" not in names:
            raise DatasetFileError("The file isn't an Excel workbook.", code="file_type_mismatch")
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            raise DatasetFileError(
                "Macro-enabled workbooks aren't supported. Save it as a regular .xlsx file.",
                code="unsupported_file_type",
            )
        if any(entry.flag_bits & 0x1 for entry in entries):
            raise DatasetFileError(
                "Encrypted workbooks aren't supported.", code="unsupported_file_type"
            )
        total = sum(entry.file_size for entry in entries)
        if total > limits.max_xlsx_uncompressed_bytes:
            raise PayloadTooLargeError(
                "The workbook expands to more data than this server accepts. Upload a smaller "
                "workbook or a CSV extract.",
                code="workbook_too_large",
            )


def _load_workbook(data: bytes, limits: IngestLimits) -> Any:
    _validate_workbook_archive(data, limits)
    # Imported lazily: openpyxl is only needed for workbooks.
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, KeyError, ValueError, zipfile.BadZipFile, OSError) as exc:
        raise DatasetFileError(
            "The workbook couldn't be read. Open it in Excel and save it again as .xlsx.",
            code="malformed_file",
        ) from exc


def list_sheets(data: bytes, limits: IngestLimits) -> list[str]:
    workbook = _load_workbook(data, limits)
    try:
        return [sheet.title for sheet in workbook.worksheets]
    finally:
        workbook.close()


def _excel_cell(value: Any) -> Any:
    # Times and durations have no column type of their own here; keep their text form.
    if isinstance(value, dt.time | dt.timedelta):
        return str(value)
    return value


def _is_empty_row(values: list[Any]) -> bool:
    return all(cell is None or (isinstance(cell, str) and not cell.strip()) for cell in values)


def read_xlsx(data: bytes, sheet_name: str | None, limits: IngestLimits) -> RawTable:
    workbook = _load_workbook(data, limits)
    try:
        sheets = [sheet.title for sheet in workbook.worksheets]
        if not sheets:
            raise DatasetFileError("The workbook has no worksheets.", code="empty_file")
        if sheet_name is None:
            if len(sheets) > 1:
                raise DatasetFileError(
                    "This workbook has several sheets. Choose the sheet to import.",
                    code="sheet_required",
                    details={"sheets": sheets},
                )
            sheet_name = sheets[0]
        elif sheet_name not in sheets:
            raise DatasetFileError(
                f'The workbook has no sheet named "{sheet_name}".',
                code="sheet_not_found",
                details={"sheets": sheets},
            )
        worksheet = workbook[sheet_name]

        # The header is the first non-empty row. It defines the table width; values in columns
        # without a header are ignored.
        header_row_index = None
        headers: list[Any] = []
        for index, row in enumerate(
            worksheet.iter_rows(values_only=True, max_col=_EXCEL_MAX_COLUMNS), start=1
        ):
            if any(cell is not None and str(cell).strip() for cell in row):
                header_row_index = index
                headers = list(row)
                break
            if index >= _MAX_LEADING_EMPTY_ROWS:
                break
        if header_row_index is None:
            raise DatasetFileError("The sheet is empty.", code="empty_file")
        while headers and (headers[-1] is None or not str(headers[-1]).strip()):
            headers.pop()
        width = len(headers)
        _check_shape(0, width, limits)

        rows: list[list[Any]] = []
        scanned = 0
        scan_limit = min(limits.max_rows + _EMPTY_ROW_ALLOWANCE, _EXCEL_MAX_ROWS)
        for row in worksheet.iter_rows(
            min_row=header_row_index + 1, max_col=width, values_only=True
        ):
            scanned += 1
            if scanned > scan_limit:
                raise DatasetFileError(
                    f"The sheet has more than {limits.max_rows:,} rows (counting empty rows). "
                    "Upload a smaller extract.",
                    code="too_many_rows",
                )
            values = [_excel_cell(cell) for cell in row]
            # Fully empty rows are usually formatting; skip them (blank CSV lines are skipped too).
            if _is_empty_row(values):
                continue
            values.extend([None] * (width - len(values)))
            rows.append(values)
            if len(rows) > limits.max_rows:
                break
        _check_shape(len(rows), width, limits)
        if not rows:
            raise DatasetFileError(
                "The sheet has column names but no data rows.", code="no_data_rows"
            )
        frame = pd.DataFrame(rows, columns=range(width), dtype=object)
        return RawTable(
            headers=["" if h is None else str(h) for h in headers],
            frame=frame,
            file_format="xlsx",
            sheet_name=sheet_name,
        )
    finally:
        workbook.close()


def parse_upload(
    filename: str, data: bytes, sheet_name: str | None, limits: IngestLimits
) -> RawTable:
    file_format = detect_format(filename, data[:16])
    if file_format == "csv":
        return read_csv(data, limits)
    return read_xlsx(data, sheet_name, limits)
